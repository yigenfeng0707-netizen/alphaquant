"""Generate 12-page AlphaQuant pitch deck. Numbers match offline case1 API / UI."""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "pitch" / "AlphaQuant-路演PPT-20260910.pptx"

NAVY = RGBColor(0x0B, 0x12, 0x20)
NAVY2 = RGBColor(0x14, 0x1E, 0x32)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
CREAM = RGBColor(0xF5, 0xEF, 0xE0)
MUTED = RGBColor(0x9A, 0xAB, 0xC4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xE0, 0x6C, 0x75)
GREEN = RGBColor(0x7D, 0xC3, 0x9A)

W = Inches(13.333)
H = Inches(7.5)
FONT = "微软雅黑"


def _set_ea(run, name: str = FONT) -> None:
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def add_run(p, text, size=18, bold=False, color=CREAM, font=FONT):
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    _set_ea(run, font)
    return run


def box(slide, l, t, w, h, fill=None, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.line.fill.background()
    if fill is not None:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    else:
        sh.fill.background()
    if line is not None:
        sh.line.fill.solid()
        sh.line.color.rgb = line
    return sh


def tb(slide, l, t, w, h):
    return slide.shapes.add_textbox(l, t, w, h)


def write_lines(tf, lines, size=18, color=CREAM, bold=False, align=PP_ALIGN.LEFT, space=10):
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        add_run(p, line, size=size, bold=bold, color=color)


def paint_bg(slide):
    box(slide, 0, 0, W, H, NAVY)
    box(slide, 0, 0, Inches(0.12), H, GOLD)
    box(slide, 0, H - Inches(0.42), W, Inches(0.42), NAVY2)
    foot = tb(slide, Inches(0.4), H - Inches(0.4), Inches(12.4), Inches(0.38))
    write_lines(
        foot.text_frame,
        ["AlphaQuant  ·  演示非实盘  ·  离线样本内回测 ≠ 收益承诺  ·  2026-09-10"],
        size=11,
        color=MUTED,
        space=0,
    )


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    paint_bg(s)
    return s


def title_bar(slide, title, subtitle=None):
    t = tb(slide, Inches(0.45), Inches(0.22), Inches(12.4), Inches(0.55))
    write_lines(t.text_frame, [title], size=28, bold=True, color=GOLD, space=0)
    if subtitle:
        s = tb(slide, Inches(0.45), Inches(0.72), Inches(12.4), Inches(0.4))
        write_lines(s.text_frame, [subtitle], size=14, color=MUTED, space=0)


def card(slide, l, t, w, h, heading, body_lines, accent=GOLD):
    box(slide, l, t, w, h, NAVY2)
    box(slide, l, t, Inches(0.08), h, accent)
    ht = tb(slide, l + Inches(0.22), t + Inches(0.12), w - Inches(0.35), Inches(0.4))
    write_lines(ht.text_frame, [heading], size=16, bold=True, color=GOLD, space=0)
    bt = tb(slide, l + Inches(0.22), t + Inches(0.52), w - Inches(0.35), h - Inches(0.65))
    write_lines(bt.text_frame, body_lines, size=14, color=CREAM, space=6)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, W, H, NAVY)
    box(s, 0, 0, W, Inches(0.18), GOLD)
    box(s, 0, H - Inches(0.18), W, Inches(0.18), GOLD)
    t = tb(s, Inches(0.8), Inches(1.4), Inches(11.6), Inches(0.5))
    write_lines(t.text_frame, ["2026 上海（长三角）中青年工程师创新创业大赛 · 金融科技"], size=16, color=GOLD)
    t = tb(s, Inches(0.8), Inches(2.0), Inches(11.6), Inches(1.2))
    write_lines(t.text_frame, ["AlphaQuant"], size=54, bold=True, color=WHITE)
    t = tb(s, Inches(0.8), Inches(3.2), Inches(11.6), Inches(0.8))
    write_lines(t.text_frame, ["基于因子增强与风险平价的智能资产配置平台"], size=22, color=CREAM)
    t = tb(s, Inches(0.8), Inches(4.3), Inches(11.6), Inches(1.6))
    write_lines(
        t.text_frame,
        [
            "个人参赛  ·  负责人 / 技术负责人 冯亦根  ·  中国电信杭州分公司（业余研发，非职务产品）",
            "可运行演示：浏览器走完 选股 → 回测 → 组合 → 风控 → 适当性 → 案例2",
            "口径：离线/回测演示，不是实盘交易，也没有接入大模型",
        ],
        size=16,
        color=MUTED,
        space=8,
    )

    # 2 痛点
    s = new_slide(prs)
    title_bar(s, "痛点：专业配置进不了浏览器", "真实需求是「能点开的研究闭环」，不是全国股民都已是用户")
    card(s, Inches(0.45), Inches(1.25), Inches(6.0), Inches(2.35), "门槛高", ["专业终端年费高，小型团队和个人用不起。", "量化流程往往要求会写代码，研究留痕靠表格拼接。"])
    card(s, Inches(6.7), Inches(1.25), Inches(6.0), Inches(2.35), "风控弱", ["常见工具停在单一收益指标。", "VaR / 压力测试、适当性匹配、归因对账不在同一套引擎。"])
    card(s, Inches(0.45), Inches(3.8), Inches(6.0), Inches(2.55), "案例 1 场景", ["个人投资者 10 万元闲置资金。", "需要：选股 → 风险平价 → 看清回撤与 VaR。", "现场按钮：仪表盘「一键演示案例 1」。"])
    card(s, Inches(6.7), Inches(3.8), Inches(6.0), Inches(2.55), "案例 2 场景", ["小型私募研究团队要可复现流程。", "批量回测 + Brinson + 标准化报告。", "1 亿元是名义演示本金，不是实盘资产管理规模。"])

    # 3 五创新点
    s = new_slide(prs)
    title_bar(s, "五个创新点（均已可点开）", "申报书承诺与当前仓库对齐；大模型策略仍是 2027 规划")
    items = [
        ("1 因子增强选股", ["12 个可计算因子", "截面标准化 + IC/IR 动态加权", "按钮：选股「刷新选股」"]),
        ("2 风险平价组合", ["SciPy SLSQP 等风险贡献", "可与等权对照", "按钮：组合「优化入选篮子」"]),
        ("3 全链路同一引擎", ["选股→回测→组合→风控", "FastAPI + Vue 3，无需编程", "侧栏七页同一套数据源"]),
        ("4 压力测试预警", ["历史模拟 VaR / CVaR", "极端窗口与情景冲击", "按钮：风控「评估风险平价组合」"]),
        ("5 投资适当性", ["问卷分档 + 匹配拦截", "本地审计，不采集身份证", "按钮：适当性「提交问卷并匹配」"]),
    ]
    for i, (h, body) in enumerate(items):
        x = Inches(0.35 + (i % 5) * 2.58)
        card(s, x, Inches(1.25), Inches(2.48), Inches(5.4), h, body)

    # 4 产品闭环
    s = new_slide(prs)
    title_bar(s, "产品闭环：七页都能走通", "数据源默认「离线演示（可断网）」；公开行情走自动并允许降级")
    pages = [
        ("仪表盘", "一键演示案例 1"),
        ("选股", "刷新选股"),
        ("回测", "运行回测"),
        ("组合", "优化入选篮子 / 加载 Brinson 归因"),
        ("风控", "评估风险平价组合"),
        ("适当性", "填入保守/平衡示例 → 提交问卷并匹配"),
        ("案例2", "运行案例 2 批量回测 / 生成并导出风控报告"),
    ]
    for i, (name, btn) in enumerate(pages):
        y = Inches(1.25 + i * 0.72)
        box(s, Inches(0.5), y, Inches(12.3), Inches(0.64), NAVY2)
        box(s, Inches(0.5), y, Inches(0.12), Inches(0.64), GOLD)
        n = tb(s, Inches(0.8), y + Inches(0.08), Inches(2.2), Inches(0.48))
        write_lines(n.text_frame, [name], size=18, bold=True, color=GOLD, space=0)
        b = tb(s, Inches(3.1), y + Inches(0.08), Inches(9.4), Inches(0.48))
        write_lines(b.text_frame, [f"主按钮：{btn}"], size=16, color=CREAM, space=0)

    # 5 案例1
    s = new_slide(prs)
    title_bar(s, "案例 1 现场演示（离线样本）", "点「一键演示案例 1」。数字来自 2026-09-10 本地 API，仪表盘两位小数。")
    kpis = [
        ("风险平价夏普", "0.20", GREEN),
        ("等权夏普", "-0.07", RED),
        ("平价是否优于等权", "是", GREEN),
        ("日度 95% VaR", "1.40%", RED),
    ]
    for i, (lab, val, col) in enumerate(kpis):
        x = Inches(0.45 + i * 3.2)
        box(s, x, Inches(1.25), Inches(3.05), Inches(1.7), NAVY2)
        a = tb(s, x + Inches(0.15), Inches(1.35), Inches(2.75), Inches(0.4))
        write_lines(a.text_frame, [lab], size=13, color=MUTED, space=0)
        b = tb(s, x + Inches(0.15), Inches(1.75), Inches(2.75), Inches(0.9))
        write_lines(b.text_frame, [val], size=32, bold=True, color=col, space=0)
    card(
        s,
        Inches(0.45),
        Inches(3.15),
        Inches(6.05),
        Inches(3.2),
        "10 万名义本金（样本内）",
        [
            "风险平价期末约 ¥103,741",
            "等权期末约 ¥95,976",
            "含近似交易成本 0.15%，非实盘成交。",
            "数据源：offline_synthetic_demo（合成面板，可断网）。",
            "口播必须带：这是历史模拟，不是未来收益保证。",
        ],
    )
    card(
        s,
        Inches(6.7),
        Inches(3.15),
        Inches(6.05),
        Inches(3.2),
        "现场操作",
        [
            "打开 http://localhost:3000",
            "确认左下角数据源 = 离线演示（可断网）",
            "仪表盘自动跑一次；可再点「一键演示案例 1」",
            "指向 KPI：风险平价夏普 / 等权夏普",
            "再切选股「刷新选股」看 12 因子与衰减告警",
        ],
    )

    # 6 适当性
    s = new_slide(prs)
    title_bar(s, "创新点五：适当性匹配（可拦截）", "本地演示规则，不是持牌机构适当性，不采集身份证")
    card(
        s,
        Inches(0.45),
        Inches(1.25),
        Inches(6.05),
        Inches(5.1),
        "现场按钮",
        [
            "侧栏「适当性」或仪表盘「适当性问卷」",
            "「填入保守示例」→「提交问卷并匹配」→ 对满仓风险平价通常「拦截」",
            "拦截是功能，不是故障。改点「填入平衡示例」再提交，通常可通过。",
            "另有「填入进取示例」。拟配置方法默认风险平价。",
            "审计表只留时间、分档、结论、答案摘要。",
        ],
    )
    card(
        s,
        Inches(6.7),
        Inches(1.25),
        Inches(6.05),
        Inches(5.1),
        "分档与边界",
        [
            "分档：保守 / 稳健 / 平衡 / 积极 / 进取",
            "匹配项：方法、仓位、波动约束",
            "不匹配就拦截并写本地 JSONL",
            "不做：证件号、合格投资者认定、代客下单",
            "评委若问「算合规吗」：不算，是可演示的产品规则。",
        ],
    )

    # 7 案例2
    s = new_slide(prs)
    title_bar(s, "案例 2：小型私募研究流程", "按钮：运行案例 2 批量回测 → 生成并导出风控报告")
    kpis = [
        ("配置效应", "4.65%"),
        ("选择效应", "3.16%"),
        ("交互效应", "-1.65%"),
        ("超额（对等权）", "6.15%"),
    ]
    for i, (lab, val) in enumerate(kpis):
        x = Inches(0.45 + i * 3.2)
        box(s, x, Inches(1.25), Inches(3.05), Inches(1.55), NAVY2)
        a = tb(s, x + Inches(0.15), Inches(1.32), Inches(2.75), Inches(0.35))
        write_lines(a.text_frame, [lab], size=13, color=MUTED, space=0)
        b = tb(s, x + Inches(0.15), Inches(1.7), Inches(2.75), Inches(0.85))
        write_lines(b.text_frame, [val], size=28, bold=True, color=GOLD, space=0)
    card(
        s,
        Inches(0.45),
        Inches(3.05),
        Inches(6.05),
        Inches(3.3),
        "口径必须说清",
        [
            "名义本金 100,000,000 元 = 机构场景演示刻度。",
            "不是把名义本金说成实盘资产管理规模。",
            "Brinson 数字可对账：4.65% + 3.16% + (-1.65%) ≈ 6.15%。",
            "报告写入 data/reports/，不是监管报备、不是托管估值。",
        ],
    )
    card(
        s,
        Inches(6.7),
        Inches(3.05),
        Inches(6.05),
        Inches(3.3),
        "组合页也可看归因",
        [
            "侧栏「组合」→「优化入选篮子」",
            "再点「加载 Brinson 归因」",
            "四种优化：等权 / 风险平价 / 最小方差 / 最大夏普",
            "风险平价夏普同样显示 0.20 vs 等权 -0.07",
        ],
    )

    # 8 技术栈
    s = new_slide(prs)
    title_bar(s, "技术栈：本机可复现", "后端 8010 · 前端 localhost:3000 · 一条 PowerShell 启动")
    card(s, Inches(0.45), Inches(1.25), Inches(6.05), Inches(5.1), "已交付", [
        "FastAPI + Uvicorn + Pydantic",
        "Vue 3 + Vite + Element Plus + ECharts",
        "Pandas / NumPy / SciPy（风险平价 SLSQP）",
        "数据：AKShare 自动；失败或断网降级 data/offline/",
        "OpenAPI：http://127.0.0.1:8010/docs",
        "启动：.\\scripts\\dev.ps1",
    ])
    card(s, Inches(6.7), Inches(1.25), Inches(6.05), Inches(5.1), "明确未接", [
        "未接入大模型（申报书第三阶段，2027-01 至 06，规划）",
        "未接支付 / SaaS 计费",
        "未接下单与券商柜台",
        "TensorFlow 不在当前选股引擎里",
        "公网魔搭 Demo 本轮不部署，避免消耗 Token",
        "本场品牌只有 AlphaQuant，不讲其他杯赛数字",
    ])

    # 9 商业模式
    s = new_slide(prs)
    title_bar(s, "商业模式（规划，产品未收费）", "评审问「怎么赚钱」：讲路径；问「现在赚多少」：还没有")
    card(s, Inches(0.45), Inches(1.25), Inches(4.0), Inches(5.1), "用户", [
        "个人投资者（案例 1）",
        "管理规模 1–10 亿的小型私募研究团队（案例 2 场景）",
        "财富管理机构的研究助理",
        "当前形态：研发中的可演示产品",
    ])
    card(s, Inches(4.65), Inches(1.25), Inches(4.0), Inches(5.1), "设想定价", [
        "专业版订阅（申报书扫描件曾写 99 元/月）",
        "机构研究席位",
        "都是设想，产品未接支付",
        "不把 TAM 说成已经占领的份额",
    ])
    card(s, Inches(8.85), Inches(1.25), Inches(4.0), Inches(5.1), "第四阶段目标", [
        "2027 下半年用户运营",
        "规划：10 万注册、付费与收入数字",
        "状态：规划，不是现状",
        "不要把规划注册数讲成已经达成",
    ])

    # 10 社会效益
    s = new_slide(prs)
    title_bar(s, "社会效益：把研究流程做成可验证的公共能力", "金融科技赛道：应用研究 + 产品研发，而不是口号普惠")
    card(s, Inches(0.45), Inches(1.25), Inches(6.05), Inches(2.4), "降低门槛", [
        "浏览器即可走完选股、配置、风控，不必先会写策略代码。",
        "离线样本保证评委电脑无网也能复核。",
    ])
    card(s, Inches(6.7), Inches(1.25), Inches(6.05), Inches(2.4), "可对账", [
        "Brinson 配置/选择/交互可加总核对超额。",
        "标准化风控报告结构化导出，研究留痕。",
    ])
    card(s, Inches(0.45), Inches(3.85), Inches(6.05), Inches(2.4), "适当性意识", [
        "用可拦截的匹配规则提醒「产品风险要和人匹配」。",
        "诚实边界：这不是持牌合规系统。",
    ])
    card(s, Inches(6.7), Inches(3.85), Inches(6.05), Inches(2.4), "工程师创新", [
        "在职工程师业余把申报书里的引擎做成可点产品。",
        "符合「正处于研发过程中的创新型项目」。",
    ])

    # 11 下一步
    s = new_slide(prs)
    title_bar(s, "下一步与诚实边界", "初赛要可验证产品；决赛再谈合作与迭代")
    card(s, Inches(0.45), Inches(1.25), Inches(6.05), Inches(5.1), "规划节奏", [
        "2026-07 至 09：核心引擎与 A 股演示（当前）",
        "2026-10 至 12：基金/债券等多资产（未开始）",
        "2027-01 至 06：探索大模型自然语言策略（规划）",
        "2027-07 至 12：用户运营目标（规划）",
        "本周：创始人电话核实黄浦初赛形式与材料清单",
    ])
    card(s, Inches(6.7), Inches(1.25), Inches(6.05), Inches(5.1), "当场禁语", [
            "其他杯赛品牌、他赛回测数字与编号（本场只用 AlphaQuant）",
            "把规划中的模型能力或注册数讲成已经上线",
            "真实持仓、已经下单、实盘成交",
            "把名义一亿元说成实盘资产管理规模",
        "把离线合成样本叫做真实行情",
        "把回测夏普 0.20 说成承诺收益",
    ])

    # 12 结束
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, W, H, NAVY)
    box(s, 0, 0, W, Inches(0.18), GOLD)
    t = tb(s, Inches(0.8), Inches(1.6), Inches(11.6), Inches(0.8))
    write_lines(t.text_frame, ["谢谢"], size=48, bold=True, color=GOLD)
    t = tb(s, Inches(0.8), Inches(2.5), Inches(11.6), Inches(3.4))
    write_lines(
        t.text_frame,
        [
            "AlphaQuant  ·  让选股、配置、风控、适当性在同一浏览器里跑完",
            "本地演示：http://localhost:3000    后端：http://127.0.0.1:8010/docs",
            "公网 Demo URL：待创始人按黄浦要求填写（本轮未部署魔搭）",
            "负责人：冯亦根    电话：18969081266    邮箱：fengyigen@qq.com",
            "黄浦初赛核实：尹老师 / 顾老师  021-33134800-21291 / 21287",
            "赛事动态：https://www.rsj.sh.gov.cn/zqncyds/",
        ],
        size=18,
        color=CREAM,
        space=12,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"PPTX {path} ({path.stat().st_size} bytes)")
